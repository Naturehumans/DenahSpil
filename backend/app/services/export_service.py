import io
import csv
import json
from typing import List, Optional
from datetime import datetime

from reportlab.lib.pagesizes import letter, landscape, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import mm
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.equipment import Equipment


# ─────────────────────────────────────────────────────────────────────────────
#  Helpers
# ─────────────────────────────────────────────────────────────────────────────

HISTORY_HEADERS = [
    "Lokasi", "Lantai", "Ruang", "ID Tempat (Slot)",
    "ID Barang (Aset)", "Kategori", "Merk", "Tipe/Model",
    "Kondisi", "Tanggal Update"
]

def _get_log_row(log) -> list:
    """Ambil nilai kolom dari satu log entry."""
    cat_name = (log.category.name if log.category else None) or "-"
    status = log.status or log.action_type or "-"
    dt = log.created_at
    tanggal = dt.strftime("%d/%m/%Y %H:%M") if dt else "-"
    return [
        log.building_name or "-",
        log.floor_name or "-",
        log.room_name or "-",
        log.slot_code or "-",
        log.asset_id or "-",
        cat_name,
        log.brand or "-",
        log.model_number or "-",
        status,
        tanggal,
    ]


# ─────────────────────────────────────────────────────────────────────────────
#  CSV
# ─────────────────────────────────────────────────────────────────────────────

def generate_history_csv(logs) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(HISTORY_HEADERS)
    for log in logs:
        writer.writerow(_get_log_row(log))
    return output.getvalue().encode("utf-8-sig")  # utf-8-sig agar Excel terbaca benar


# ─────────────────────────────────────────────────────────────────────────────
#  XLSX
# ─────────────────────────────────────────────────────────────────────────────

def generate_history_xlsx(logs, title: str = "Riwayat Data Barang") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Riwayat"

    # ── Warna / Style ──────────────────────────────────────────────
    GREEN = "3A9542"
    LIGHT_GREEN = "E8F5E9"
    LIGHT_GRAY = "F8FAFC"
    WHITE = "FFFFFF"
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=GREEN)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    title_font = Font(name="Calibri", bold=True, size=14, color="1B5E20")
    subtitle_font = Font(name="Calibri", size=10, color="64748B", italic=True)

    # ── Judul ──────────────────────────────────────────────────────
    ws.merge_cells("A1:J1")
    ws["A1"] = title
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:J2")
    ws["A2"] = f"Diekspor pada: {datetime.now().strftime('%d %B %Y, %H:%M')}"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    ws.append([])  # baris kosong
    ws.row_dimensions[3].height = 6

    # ── Header tabel ──────────────────────────────────────────────
    header_row = 4
    for col_idx, h in enumerate(HISTORY_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    ws.row_dimensions[header_row].height = 28

    # ── Data ──────────────────────────────────────────────────────
    for row_idx, log in enumerate(logs, start=1):
        excel_row = header_row + row_idx
        row_data = _get_log_row(log)
        is_even = row_idx % 2 == 0
        row_fill = PatternFill("solid", fgColor=LIGHT_GREEN if is_even else WHITE)

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=excel_row, column=col_idx, value=val)
            cell.fill = row_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
            cell.font = Font(name="Calibri", size=10)
        ws.row_dimensions[excel_row].height = 18

    # ── Lebar kolom ───────────────────────────────────────────────
    col_widths = [18, 12, 16, 18, 22, 14, 14, 16, 18, 18]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Freeze header ─────────────────────────────────────────────
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
#  PDF
# ─────────────────────────────────────────────────────────────────────────────

def generate_history_pdf(logs, title: str = "Riwayat Data Barang") -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "title",
        parent=styles["Heading1"],
        fontSize=16,
        textColor=colors.HexColor("#1B5E20"),
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        "subtitle",
        parent=styles["Normal"],
        fontSize=9,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=12,
    )
    cell_style = ParagraphStyle(
        "cell",
        parent=styles["Normal"],
        fontSize=7.5,
        leading=10,
    )

    elements = [
        Paragraph(title, title_style),
        Paragraph(
            f"Diekspor pada: {datetime.now().strftime('%d %B %Y, %H:%M')}  |  Total data: {len(logs)} baris",
            subtitle_style,
        ),
    ]

    # ── Table data ─────────────────────────────────────────────────
    header_cells = [Paragraph(f"<b>{h}</b>", cell_style) for h in HISTORY_HEADERS]
    data = [header_cells]
    for log in logs:
        row = [Paragraph(str(v), cell_style) for v in _get_log_row(log)]
        data.append(row)

    col_widths_pt = [55, 38, 48, 55, 65, 42, 38, 48, 52, 52]  # points, total ~493
    table = Table(data, colWidths=col_widths_pt, repeatRows=1)

    GREEN_COLOR = colors.HexColor("#3A9542")
    LIGHT_GREEN = colors.HexColor("#E8F5E9")
    LIGHT_GRAY = colors.HexColor("#F8FAFC")

    style = TableStyle([
        # Header
        ("BACKGROUND", (0, 0), (-1, 0), GREEN_COLOR),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        # Alternating row colors
        *[
            ("BACKGROUND", (0, i), (-1, i), LIGHT_GREEN if i % 2 == 0 else LIGHT_GRAY)
            for i in range(1, len(data))
        ],
        # Grid
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 1), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 4),
    ])
    table.setStyle(style)
    elements.append(table)

    doc.build(elements)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
#  Legacy Equipment Export (dipertahankan agar endpoint lama tidak rusak)
# ─────────────────────────────────────────────────────────────────────────────

def generate_csv(equipments: List[Equipment]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Name", "Category", "Floor", "Brand", "Model",
        "Installation Date", "Lifespan (Months)", "Expiry Date", "Status", "Notes"
    ])
    for eq in equipments:
        writer.writerow([
            eq.name,
            eq.category.name if eq.category else "-",
            eq.floor.name if eq.floor else "-",
            eq.brand or "-",
            eq.model_number or "-",
            eq.installation_date.isoformat() if eq.installation_date else "-",
            eq.lifespan_months,
            eq.expiry_date.isoformat() if eq.expiry_date else "-",
            eq.status,
            eq.notes or "-"
        ])
    return output.getvalue().encode("utf-8")


def generate_pdf(equipments: List[Equipment], title: str = "Equipment Report") -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter))
    elements = []
    styles = getSampleStyleSheet()
    elements.append(Paragraph(title, styles["Heading1"]))
    elements.append(Paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles["Normal"]))
    elements.append(Spacer(1, 12))
    data = [["Name", "Category", "Floor", "Brand", "Installation Date", "Expiry Date", "Status"]]
    for eq in equipments:
        data.append([
            eq.name,
            eq.category.name if eq.category else "-",
            eq.floor.name if eq.floor else "-",
            eq.brand or "-",
            eq.installation_date.strftime("%Y-%m-%d") if eq.installation_date else "-",
            eq.expiry_date.strftime("%Y-%m-%d") if eq.expiry_date else "-",
            eq.status
        ])
    table = Table(data)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3a9542")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
        ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#f0f3f7")),
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
    ]))
    elements.append(table)
    doc.build(elements)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes
